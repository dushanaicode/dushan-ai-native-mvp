import asyncio
import io
import json
import os
import struct
import subprocess
import sys
import threading
import zlib
from lzma import LZMAError
from pathlib import Path
from types import SimpleNamespace
from typing import Annotated
from zipfile import ZIP_DEFLATED, ZIP_LZMA, ZipFile

import pytest
from openpyxl import Workbook
from pydantic import BaseModel, field_validator

from fixtures.config_factory import ConfigFactory
from framework.starter_excel.config.excel_settings import ExcelSettings
from framework.starter_excel.core.excel_schema import ExcelSchema
from framework.starter_excel.definitions.constants.excel_error_codes import ExcelErrorCodes
from framework.starter_excel.exception.excel_exception import ExcelException
from framework.starter_excel.handler.excel_upload_validator import ExcelUploadValidator
from framework.starter_excel.model.excel_column import ExcelColumn
from framework.starter_excel.reader import excel_reader as reader_module
from framework.starter_excel.reader.excel_reader import ExcelReader
from framework.starter_excel.writer.excel_writer import ExcelWriter


class TextRow(BaseModel):
    text: Annotated[str, ExcelColumn("文本")]


@pytest.mark.parametrize("fields", [["missing"], ["text", "text"], []])
def test_export_field_selection_errors_are_request_validation(fields):
    with pytest.raises(ExcelException) as caught:
        ExcelSchema(TextRow).export_columns(fields)
    assert caught.value.error_code is ExcelErrorCodes.VALIDATION


def test_model_without_exportable_columns_remains_configuration_error():
    class Sensitive(BaseModel):
        token: Annotated[str, ExcelColumn("令牌")]

    with pytest.raises(ExcelException) as caught:
        ExcelSchema(Sensitive).export_columns()
    assert caught.value.error_code is ExcelErrorCodes.CONFIG


@pytest.mark.parametrize("failure", [None, "limit", "control", "structure"])
def test_failed_rewind_preserves_primary_or_reports_read_failure(monkeypatch, failure):
    class BrokenRewind(io.BytesIO):
        broken = False

        def seek(self, offset, whence=io.SEEK_SET):
            if self.broken:
                raise OSError("rewind unavailable: private-stream-detail")
            return super().seek(offset, whence)

    source = upload()
    content = source.file.getvalue()
    source.file.close()
    stream = BrokenRewind(content)
    source.file = stream
    validator = ExcelUploadValidator(settings())
    original = validator._archive
    primary = None
    if failure == "limit":
        primary = ExcelException(ExcelErrorCodes.LIMIT, "校验上限")
    elif failure == "control":
        primary = BaseException("control exit")
    elif failure == "structure":
        primary = ValueError("invalid archive")

    def archive(value):
        original(value)
        stream.broken = True
        if primary is not None:
            raise primary

    monkeypatch.setattr(validator, "_archive", archive)
    try:
        with pytest.raises(BaseException if failure == "control" else ExcelException) as caught:
            validator.validate(source)
        if primary is None:
            assert caught.value.error_code is ExcelErrorCodes.READ
            assert isinstance(caught.value.__cause__, OSError)
        elif failure == "structure":
            assert caught.value.error_code is ExcelErrorCodes.VALIDATION
            assert caught.value.__cause__ is primary
        else:
            assert caught.value is primary
        if primary is not None:
            assert caught.value.__notes__ == ["XLSX 上传流位置恢复失败：OSError"]
        assert not stream.closed
    finally:
        stream.close()


def settings(**changes):
    values = ConfigFactory.values()["config"]["models"]["excel"]
    values.update(changes)
    return ExcelSettings.model_validate(values)


def upload():
    workbook = Workbook()
    output = io.BytesIO()
    try:
        workbook.active.title = "活动表"
        workbook.active.append(["文本"])
        workbook.active.append(["active"])
        sheet = workbook.create_sheet("指定表")
        sheet.append(["文本"])
        sheet.append(["selected"])
        workbook.save(output)
        output.seek(7)
        return SimpleNamespace(file=output, filename="book.xlsx", content_type=None)
    finally:
        workbook.close()


@pytest.mark.parametrize("entrypoint", ["validator", "reader"])
@pytest.mark.parametrize(
    ("corruption", "compression", "cause_type"),
    [
        ("deflate", ZIP_DEFLATED, zlib.error),
        ("lzma", ZIP_LZMA, LZMAError),
        ("unsupported", ZIP_DEFLATED, NotImplementedError),
    ],
)
async def test_malformed_zip_compression_is_validation_error(
    entrypoint, corruption, compression, cause_type
):
    source = upload()
    try:
        with ZipFile(source.file) as original, io.BytesIO() as output:
            with ZipFile(output, "w", compression=compression) as archive:
                for entry in original.infolist():
                    archive.writestr(entry.filename, original.read(entry))
            payload = bytearray(output.getvalue())
    finally:
        source.file.close()
    offset = 30 + sum(struct.unpack_from("<HH", payload, 26))
    if corruption == "deflate":
        payload[offset] = 0x07  # DEFLATE 保留块类型，实际解压必须失败。
    elif corruption == "lzma":
        payload[offset + 4] = 0xFF  # ZIP LZMA 头之后的非法过滤器属性。
    else:
        central = payload.find(b"PK\x01\x02")
        struct.pack_into("<H", payload, 8, 99)
        struct.pack_into("<H", payload, central + 10, 99)
    with io.BytesIO(payload) as stream:
        source.file = stream
        stream.seek(7)
        with pytest.raises(ExcelException) as caught:
            if entrypoint == "validator":
                ExcelUploadValidator(settings()).validate(source)
            else:
                await ExcelReader(settings()).read(source, TextRow)
        assert caught.value.error_code is ExcelErrorCodes.VALIDATION
        assert isinstance(caught.value.__cause__, cause_type)
        assert stream.tell() == 7
        assert not stream.closed


async def wait_event(event):
    assert await asyncio.to_thread(event.wait, 5), "受测线程没有到达预期边界"


async def cancel_twice(task):
    task.cancel("first")
    await asyncio.sleep(0.01)
    task.cancel("second")
    await asyncio.sleep(0.01)
    assert not task.done(), "取消不应提前释放仍在使用文件的线程"


async def test_named_sheet_and_missing_sheet_restore_borrowed_stream():
    reader = ExcelReader(settings())
    source = upload()
    assert await reader.read(source, TextRow) == [TextRow(text="active")]
    assert await reader.read(source, TextRow, sheet_name="指定表") == [TextRow(text="selected")]
    with pytest.raises(ExcelException) as caught:
        await reader.read(source, TextRow, sheet_name="不存在")
    assert caught.value.error_code == ExcelErrorCodes.VALIDATION
    assert "工作表不存在" in caught.value.msg
    assert source.file.tell() == 7 and not source.file.closed
    source.file.close()


@pytest.mark.parametrize("phase", ["validation", "load", "rows"])
async def test_repeated_cancel_waits_for_file_worker_before_cleanup(monkeypatch, phase):
    source = upload()
    started, release, finished = (threading.Event() for _ in range(3))
    workbooks = []
    original_load = reader_module.load_workbook
    original_validate = ExcelUploadValidator.validate
    original_islice = reader_module.islice

    def pause():
        started.set()
        assert release.wait(5)
        assert not source.file.closed
        finished.set()

    def load(*args, **kwargs):
        workbook = original_load(*args, **kwargs)
        workbooks.append(workbook)
        if phase == "load":
            pause()
        return workbook

    def validate(self, value):
        if phase == "validation":
            value.file.seek(17)
            pause()
        return original_validate(self, value)

    def batch(rows, count):
        if phase == "rows":
            pause()
        yield from original_islice(rows, count)

    monkeypatch.setattr(reader_module, "load_workbook", load)
    monkeypatch.setattr(ExcelUploadValidator, "validate", validate)
    monkeypatch.setattr(reader_module, "islice", batch)
    task = asyncio.create_task(ExcelReader(settings()).read(source, TextRow))
    try:
        await wait_event(started)
        await cancel_twice(task)
        assert not source.file.closed
        if phase == "validation":
            assert source.file.tell() == 17
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError) as caught:
        await task
    assert caught.value.args == ("first",)
    assert finished.is_set()
    assert all(workbook._archive.fp is None for workbook in workbooks)
    assert source.file.tell() == 7 and not source.file.closed
    source.file.close()


async def test_cancel_and_worker_failure_preserve_both_causes(monkeypatch):
    started, release = threading.Event(), threading.Event()
    source = upload()
    failure = RuntimeError("worker parse failure")

    def validate(self, value):
        value.file.seek(17)
        started.set()
        assert release.wait(5)
        assert not value.file.closed
        raise failure

    monkeypatch.setattr(ExcelUploadValidator, "validate", validate)
    task = asyncio.create_task(ExcelReader(settings()).read(source, TextRow))
    try:
        await wait_event(started)
        await cancel_twice(task)
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError) as caught:
        await task
    assert caught.value.args == ("first",) and caught.value.__cause__ is failure
    assert source.file.tell() == 7 and not source.file.closed
    source.file.close()


async def test_repeated_cancel_during_close_waits_and_keeps_close_error(monkeypatch):
    started, release = threading.Event(), threading.Event()
    source = upload()
    failure = OSError("close failure")
    original_load = reader_module.load_workbook
    original_close = Workbook.close

    def close(workbook):
        started.set()
        assert release.wait(5)
        original_close(workbook)
        raise failure

    def load(*args, **kwargs):
        workbook = original_load(*args, **kwargs)
        workbook.close = lambda: close(workbook)
        return workbook

    monkeypatch.setattr(reader_module, "load_workbook", load)
    task = asyncio.create_task(ExcelReader(settings()).read(source, TextRow))
    try:
        await wait_event(started)
        await cancel_twice(task)
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError) as caught:
        await task
    assert caught.value.args == ("first",)
    assert caught.value.__cause__.exceptions[0].exceptions == (failure,)
    assert source.file.tell() == 7 and not source.file.closed
    source.file.close()


@pytest.mark.parametrize("limit", [1, 2])
async def test_import_concurrency_gate_covers_worker_until_completion(monkeypatch, limit):
    entered, release = threading.Event(), threading.Event()
    lock = threading.Lock()
    starts = 0
    original = ExcelReader._load_workbook

    def load(self, source):
        nonlocal starts
        with lock:
            starts += 1
            if starts == limit:
                entered.set()
        assert release.wait(5)
        return original(self, source)

    monkeypatch.setattr(ExcelReader, "_load_workbook", load)
    reader = ExcelReader(settings(max_concurrent_imports=limit))
    sources = [upload() for _ in range(limit + 1)]
    tasks = [asyncio.create_task(reader.read(source, TextRow)) for source in sources]
    try:
        await wait_event(entered)
        await asyncio.sleep(0.02)
        assert starts == limit
        tasks[-1].cancel("queued")
        with pytest.raises(asyncio.CancelledError):
            await tasks[-1]
        assert starts == limit and sources[-1].file.tell() == 7
    finally:
        release.set()
    assert await asyncio.gather(*tasks[:-1]) == [[TextRow(text="active")]] * limit
    for source in sources:
        assert not source.file.closed
        source.file.close()


async def test_business_converter_and_model_callbacks_stay_on_application_thread():
    application_thread = threading.get_ident()
    callbacks = []

    class Converter:
        async def to_excel(self, value, context):
            assert threading.get_ident() == application_thread
            callbacks.append("export")
            return value

        async def to_python(self, value, context):
            assert threading.get_ident() == application_thread
            callbacks.append("import")
            return value

    class Row(BaseModel):
        text: Annotated[str, ExcelColumn("文本", converter=Converter())]

        @field_validator("text")
        @classmethod
        def check(cls, value):
            assert threading.get_ident() == application_thread
            callbacks.append("model")
            return value

    source = upload()
    rows = await ExcelReader(settings()).read(source, Row)
    with await ExcelWriter(settings()).write("数据", Row, rows) as output:
        assert output.getbuffer().nbytes > 0
    assert callbacks == ["import", "model", "export"]
    source.file.close()


@pytest.mark.parametrize("worker_fails", [False, True])
async def test_save_cancellation_waits_before_closing_workbook_and_output(
    monkeypatch, worker_fails
):
    started, release, finished = (threading.Event() for _ in range(3))
    original_save = Workbook.save
    original_close = Workbook.close
    buffers, workbooks = [], []
    failure = OSError("save failure")

    def save(workbook, output):
        buffers.append(output)
        workbooks.append(workbook)
        workbook.was_closed = False
        output.write(b"worker-started")
        started.set()
        assert release.wait(5)
        assert not output.closed and not workbook.was_closed
        finished.set()
        if worker_fails:
            raise failure
        original_save(workbook, output)

    def close(workbook):
        workbook.was_closed = True
        original_close(workbook)

    monkeypatch.setattr(Workbook, "save", save)
    monkeypatch.setattr(Workbook, "close", close)
    task = asyncio.create_task(ExcelWriter(settings()).write("数据", TextRow, [TextRow(text="x")]))
    try:
        await wait_event(started)
        await cancel_twice(task)
        assert not buffers[0].closed and not workbooks[0].was_closed
    finally:
        release.set()
    with pytest.raises(asyncio.CancelledError) as caught:
        await task
    assert caught.value.args == ("first",) and finished.is_set()
    assert buffers[0].closed and workbooks[0].was_closed
    if worker_fails:
        assert caught.value.__cause__ is failure


def test_vba_resource_and_physical_encryption_flags_are_rejected():
    source = upload()
    content = source.file.getvalue()
    source.file.close()
    macro = io.BytesIO(content)
    with ZipFile(macro, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"test-vba-payload")
    with pytest.raises(ExcelException) as caught:
        ExcelUploadValidator(settings()).validate(
            SimpleNamespace(file=macro, filename="macro.xlsx", content_type=None)
        )
    assert isinstance(caught.value.__cause__, ValueError)
    assert "VBA" in str(caught.value.__cause__)
    assert not macro.closed
    encrypted = bytearray(content)
    central = encrypted.find(b"PK\x01\x02")
    # ZipFile.writestr 会清除 flag_bits；须在已完成的 ZIP 字节中同时置位两处标志。
    struct.pack_into("<H", encrypted, 6, struct.unpack_from("<H", encrypted, 6)[0] | 1)
    struct.pack_into(
        "<H", encrypted, central + 8, struct.unpack_from("<H", encrypted, central + 8)[0] | 1
    )
    assert struct.unpack_from("<H", encrypted, 6)[0] & 1
    assert struct.unpack_from("<H", encrypted, central + 8)[0] & 1
    with pytest.raises(ExcelException) as caught:
        ExcelUploadValidator(settings()).validate(
            SimpleNamespace(
                file=io.BytesIO(encrypted), filename="encrypted.xlsx", content_type=None
            )
        )
    assert isinstance(caught.value.__cause__, ValueError)
    assert "加密" in str(caught.value.__cause__)


def test_every_excel_module_and_plain_roundtrip_work_when_ip_import_is_blocked(tmp_path):
    root = Path(__file__).resolve().parents[3]
    config = tmp_path / "settings.json"
    config.write_text(json.dumps(settings().model_dump(), ensure_ascii=False), encoding="utf-8")
    script = tmp_path / "without_ip.py"
    script.write_text(
        """import asyncio
import importlib.abc
import io
import json
import pkgutil
import sys
from pathlib import Path
from types import SimpleNamespace
from typing import Annotated

sys.path.insert(0, sys.argv[1])
class RejectIp(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'framework.starter_ip' or fullname.startswith('framework.starter_ip.'):
            raise AssertionError('Excel attempted IP import: ' + fullname)

sys.meta_path.insert(0, RejectIp())
import framework.starter_excel
for info in pkgutil.walk_packages(framework.starter_excel.__path__, 'framework.starter_excel.'):
    __import__(info.name)
from pydantic import BaseModel
from framework.starter_excel.config.excel_settings import ExcelSettings
from framework.starter_excel.model.excel_column import ExcelColumn
from framework.starter_excel.reader.excel_reader import ExcelReader
from framework.starter_excel.writer.excel_writer import ExcelWriter
class Row(BaseModel):
    text: Annotated[str, ExcelColumn('text')]
async def main():
    settings = ExcelSettings.model_validate(json.loads(Path(sys.argv[2]).read_text(encoding='utf-8')))
    with await ExcelWriter(settings).write('data', Row, [Row(text='standalone')]) as stream:
        upload = SimpleNamespace(file=stream, filename='file.xlsx', content_type=None)
        assert await ExcelReader(settings).read(upload, Row) == [Row(text='standalone')]
    assert not any(name.startswith('framework.starter_ip') for name in sys.modules)
asyncio.run(main())
print('standalone Excel passed')
""",
        encoding="utf-8",
    )
    result = subprocess.run(
        [
            sys.executable,
            "-B",
            "-X",
            "utf8",
            "-I",
            str(script),
            str(root / "dushan-admin-backend"),
            str(config),
        ],
        cwd=root,
        env=os.environ.copy(),
        text=True,
        capture_output=True,
        timeout=20,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stdout.strip() == "standalone Excel passed"
