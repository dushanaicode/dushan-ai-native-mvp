from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.asyncio(loop_scope="module")


@pytest.mark.parametrize(
    "path", ["job", "mq", "data-source", "logger/api-access-log", "logger/api-error-log"]
)
async def test_selected_export_fields_match_workbook_and_reject_unknown_fields(admin_client, path):
    prefix = "/admin-api/infra/" + path
    fields = (await admin_client.get(prefix + "/export-fields")).json()
    assert fields["code"] == 0 and fields["data"], fields
    selected = fields["data"][0]
    response = await admin_client.get(
        prefix + "/export-excel", params={"fields": selected["field"]}
    )
    assert response.status_code == 200 and response.content.startswith(b"PK"), response.text[:200]
    assert ".xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert list(next(workbook.active.iter_rows(values_only=True))) == [selected["title"]]
    finally:
        workbook.close()
    invalid = (
        await admin_client.get(prefix + "/export-excel", params={"fields": "unknownPrivateField"})
    ).json()
    assert invalid["code"] not in (0, 500), invalid
    page = (await admin_client.get(prefix + "/page", params={"fields": selected["field"]})).json()
    assert page["code"] == 422, page


async def test_mq_options_only_allow_current_declarations(admin_client):
    prefix = "/admin-api/infra/mq"
    options = (await admin_client.get(prefix + "/consumers")).json()
    assert options["code"] == 0 and options["data"], options
    unknown = (
        await admin_client.post(
            prefix + "/create",
            json={"topic": "unknown" + uuid4().hex, "consumer": "not-deployed", "retryCount": 0},
        )
    ).json()
    assert unknown["code"] not in (0, 500) and "未注册" in unknown["message"], unknown
    declaration = options["data"][0]
    existing = (
        await admin_client.get(prefix + "/page", params={"consumer": declaration["key"]})
    ).json()
    assert existing["code"] == 0, existing
    payload = {
        "topic": declaration["topic"],
        "consumer": declaration["key"],
        "retryCount": declaration["retryCount"],
        "description": "declaration test",
        "enabled": True,
    }
    if existing["data"]["items"]:
        payload["id"] = existing["data"]["items"][0]["id"]
        method, path = admin_client.put, "/update"
    else:
        method, path = admin_client.post, "/create"
    invalid = (
        await method(prefix + path, json={**payload, "retryCount": payload["retryCount"] + 1})
    ).json()
    assert invalid["code"] not in (0, 500) and "声明一致" in invalid["message"], invalid
    accepted = (await method(prefix + path, json=payload)).json()
    assert accepted["code"] == 0, accepted
