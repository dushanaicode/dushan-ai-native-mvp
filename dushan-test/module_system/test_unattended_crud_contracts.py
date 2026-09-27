from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

pytestmark = pytest.mark.asyncio(loop_scope="module")

IMPORT_PASSWORD = "ImportTest123!"


@pytest.fixture(scope="module")
def system_settings_overrides():
    return {"default_password": IMPORT_PASSWORD}


def stored(connection, table, fields, identifier):
    with connection.cursor() as cursor:
        cursor.execute(f"SELECT {','.join(fields)} FROM {table} WHERE id=%s", (identifier,))
        return cursor.fetchone()


async def test_department_blank_email_round_trip(admin_client):
    body = {"name": "empty" + uuid4().hex[:8], "parentId": "0", "sort": 1, "status": 1, "email": ""}
    created = (await admin_client.post("/admin-api/system/dept/create", json=body)).json()
    assert created["code"] == 0, created
    body["id"] = created["data"]
    for email in (None, "valid@example.com", ""):
        body["email"] = email
        result = (await admin_client.put("/admin-api/system/dept/update", json=body)).json()
        assert result["code"] == 0, result
        found = (
            await admin_client.get("/admin-api/system/dept/get", params={"id": body["id"]})
        ).json()
        assert found["data"]["email"] == (email or None)
    invalid = (
        await admin_client.put("/admin-api/system/dept/update", json={**body, "email": "invalid"})
    ).json()
    assert invalid["code"] == 422


@pytest.mark.parametrize("kind", ["mail", "sms", "social", "oauth"])
async def test_edit_omitted_credentials_preserves_and_explicit_change_rotates(
    kind, admin_client, system_database
):
    suffix = uuid4().hex[:8]
    secret = "UnattendedOldSecret1234567890!AaBbCcDd"
    rotated = "UnattendedNewSecret1234567890!EeFfGgHh"
    if kind == "mail":
        path, table, keys, columns, display = (
            "mail/account",
            "system_mail_account",
            ["password"],
            ["password"],
            "username",
        )
        body = {
            "mail": "sender@example.com",
            "username": "sender",
            "password": secret,
            "host": "localhost",
            "port": 1025,
            "sslEnable": False,
            "starttlsEnable": False,
        }
    elif kind == "sms":
        path, table, keys, columns, display = (
            "sms/channel",
            "system_sms_channel",
            ["apiKey", "apiSecret"],
            ["api_key", "api_secret"],
            "remark",
        )
        body = {
            "signature": "Fixture",
            "code": "ALIYUN",
            "status": 0,
            "apiKey": secret,
            "apiSecret": secret,
            "remark": "before",
        }
    elif kind == "social":
        path, table, keys, columns, display = (
            "social/client",
            "system_social_client",
            ["clientSecret"],
            ["client_secret"],
            "name",
        )
        body = {
            "name": "fixture" + suffix,
            "socialType": 20,
            "userType": 2,
            "clientId": "fixture" + suffix,
            "clientSecret": secret,
            "status": 0,
        }
    else:
        path, table, keys, columns, display = (
            "oauth2/client",
            "system_oauth2_client",
            ["secret"],
            ["secret"],
            "name",
        )
        body = {
            "clientId": "fixture" + suffix,
            "secret": secret,
            "name": "fixture",
            "logo": "https://example.com/logo.png",
            "status": 1,
            "userType": 2,
            "accessTokenValiditySeconds": 1800,
            "refreshTokenValiditySeconds": 2592000,
            "redirectUris": ["https://example.com/callback"],
            "authorizedGrantTypes": ["authorization_code", "refresh_token"],
            "scopes": [],
            "autoApproveScopes": [],
            "authorities": [],
            "resourceIds": [],
        }
    endpoint = "/admin-api/system/" + path
    created = (await admin_client.post(endpoint + "/create", json=body)).json()
    assert created["code"] == 0, created
    identifier = created["data"]
    connection = system_database[2]
    before = stored(connection, table, columns, identifier)
    update = {key: value for key, value in body.items() if key not in keys}
    update.update(id=identifier)
    update[display] = "updated" + suffix
    result = (await admin_client.put(endpoint + "/update", json=update)).json()
    assert result["code"] == 0, result
    assert stored(connection, table, columns, identifier) == before
    visible = (await admin_client.get(endpoint + "/get", params={"id": identifier})).json()
    assert secret not in str(visible)
    result = (
        await admin_client.put(
            endpoint + "/update", json={**update, **{key: rotated for key in keys}}
        )
    ).json()
    assert result["code"] == 0, result
    assert stored(connection, table, columns, identifier) == tuple(rotated for _ in columns)
    missing = (
        await admin_client.post(
            endpoint + "/create",
            json={key: value for key, value in body.items() if key not in keys},
        )
    ).json()
    assert missing["code"] == 422, missing
    forged = (
        await admin_client.post(endpoint + "/create", json={**update, "id": "900000000000000123"})
    ).json()
    assert forged["code"] != 0


async def test_role_ids_are_strings_and_owner_role_cannot_be_transferred(
    admin_client, system_database
):
    suffix = uuid4().hex[:8]
    user = (
        await admin_client.post(
            "/admin-api/system/user/create",
            json={"username": "grant" + suffix, "nickname": "Grant", "password": "GrantTest123!"},
        )
    ).json()
    role = (
        await admin_client.post(
            "/admin-api/system/permission/role/create",
            json={"name": "Grant" + suffix, "code": "grant" + suffix, "sort": 1},
        )
    ).json()
    assert user["code"] == role["code"] == 0, (user, role)
    assigned = (
        await admin_client.post(
            "/admin-api/system/permission/assign-user-role",
            json={"userId": user["data"], "roleIds": [role["data"]]},
        )
    ).json()
    assert assigned["code"] == 0, assigned
    found = (
        await admin_client.get(
            "/admin-api/system/permission/list-user-roles", params={"userId": user["data"]}
        )
    ).json()
    assert found["data"] == [role["data"]]
    assert isinstance(found["data"][0], str) and int(found["data"][0]) > 2**53
    with system_database[2].cursor() as cursor:
        cursor.execute("SELECT id FROM system_role WHERE code='super_admin' AND deleted=0")
        owner_role_id = str(cursor.fetchone()[0])
    denied = (
        await admin_client.post(
            "/admin-api/system/permission/assign-user-role",
            json={"userId": user["data"], "roleIds": [owner_role_id]},
        )
    ).json()
    assert denied["code"] != 0
    unchanged = (
        await admin_client.get(
            "/admin-api/system/permission/list-user-roles", params={"userId": user["data"]}
        )
    ).json()
    assert unchanged["data"] == [role["data"]]


def workbook(rows):
    document = Workbook()
    sheet = document.active
    sheet.append(["用户账号", "用户昵称", "部门编号", "用户邮箱", "手机号码", "性别", "状态"])
    for row in rows:
        sheet.append(row)
    target = BytesIO()
    document.save(target)
    document.close()
    return target.getvalue()


async def test_import_template_create_duplicate_update_and_partial_failure(
    admin_client, system_database
):
    template = await admin_client.get("/admin-api/system/user/get-import-template")
    assert template.status_code == 200 and template.content.startswith(b"PK")
    document = load_workbook(BytesIO(template.content))
    assert document.active.max_row >= 2
    document.close()
    password = IMPORT_PASSWORD
    name = "import" + uuid4().hex[:8]

    async def send(rows, update):
        response = await admin_client.post(
            "/admin-api/system/user/import",
            data={"updateSupport": str(update).lower()},
            files={
                "file": (
                    "users.xlsx",
                    workbook(rows),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        result = response.json()
        assert result["code"] == 0, result
        assert password not in response.text
        return result["data"]

    row = [name, "Imported", None, None, None, "男", "开启"]
    result = await send([row, ["bad!", "Invalid", None, None, None, "男", "开启"]], False)
    assert result["createUsernames"] == [name]
    assert "bad!" in result["failureUsernames"]
    duplicate = await send([row], False)
    assert name in duplicate["failureUsernames"]
    row[1] = "Updated"
    updated = await send([row], True)
    assert updated["updateUsernames"] == [name]
    with system_database[2].cursor() as cursor:
        cursor.execute(
            "SELECT nickname,status FROM system_users WHERE username=%s AND deleted=0", (name,)
        )
        assert cursor.fetchone() == ("Updated", 1)
    login = (
        await admin_client.post(
            "/admin-api/system/auth/login",
            json={"username": name, "password": password},
            headers={"X-Tenant-Id": "1"},
        )
    ).json()
    assert login["code"] == 0, login
